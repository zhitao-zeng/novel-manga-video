import json
from pathlib import Path
from docx import Document
from novel_manga.source.ingest import read_novel, chapter_table
from novel_manga.planning.authored_brief import chapter_text
from novel_manga.application.planning.lean_bible import chapter_texts
from novel_manga.source.splitting import split_episodes, visible_count
from novel_manga.planning import parts
from novel_manga.application.planning.parts import _prompt


def metadata(novel, *, offsets=True):
    return {'source':str(novel.source_path), 'chapters':[
        {'index':e.index, 'title':e.source_title, **({'start':e.source_start,'end':e.source_end} if offsets else {})}
        for e in novel.episodes]}


def test_normalized_offsets_are_shared_by_both_readers(tmp_path):
    path=tmp_path/'book.txt'
    path.write_text('\ufeff\n\n第一章 开端\n甲说你好。\n\n\n\n\n第二章 后续\n乙回答再见。\n',encoding='utf-8')
    novel=read_novel(path,'book');meta=metadata(novel)
    texts=chapter_texts(meta,[1,2])
    for ep in novel.episodes:
        assert chapter_text(meta,ep.index).strip()==texts[ep.index]==ep.source_text


def test_docx_registered_parts_use_extracted_text(tmp_path):
    path=tmp_path/'book.docx';doc=Document()
    doc.add_paragraph('第一章 长章')
    for _ in range(90):doc.add_paragraph('情节推进'+ '甲'*90+'。')
    doc.save(path)
    novel=read_novel(path,'book');meta=metadata(novel)
    assert len(novel.episodes)==2
    assert chapter_texts(meta,[1,2])=={e.index:e.source_text for e in novel.episodes}
    assert chapter_text(meta,2).strip()==novel.episodes[1].source_text


def test_existing_registration_keeps_chapter_and_ledger_numbers(tmp_path):
    path=tmp_path/'book.txt';text='第一章 长章\n'+('甲说你好。'*1500)+'\n第二章 后续\n乙回答再见。'
    path.write_text(text)
    old,_=split_episodes(text,split_long_chapters=False)
    meta={'source':str(path),'chapters':[{'index':e.index,'title':e.source_title} for e in old]}
    (tmp_path/'novel.json').write_text(json.dumps(meta))
    registered=read_novel(path,'book',chapters=chapter_table(tmp_path))
    assert len(read_novel(path,'book').episodes)==3
    assert len(registered.episodes)==2
    assert registered.episodes[1].source_title=='第二章 后续'
    assert chapter_texts(meta,[2])[2]==registered.episodes[1].source_text
    # Reinitializing with the registered table must not change either chapter.
    assert [e.source_text for e in registered.episodes]==[e.source_text for e in old]


def test_long_chapter_beyond_ten_parts_loses_no_source():
    source='第一章 超长\n'+('甲'*99+'。')*650
    episodes,_=split_episodes(source)
    assert len(episodes)==11
    assert sum(e.text_count for e in episodes)==visible_count(source)
    assert all(source[e.source_start:e.source_end]==e.source_text for e in episodes)


def test_many_short_drama_parts_roundtrip_without_renumbering_chapter(tmp_path):
    rows=['他说：“'+ '话'*99+'。”']*21
    count=parts.suggest_count(rows)
    assert count>3
    cut=parts.mechanical_split(rows,count)
    assert parts.validate(cut,rows)==[]
    parts.write_parts(tmp_path,cut,decided_by='test',chapter=12)
    assert parts.read_parts(tmp_path)==cut
    assert json.loads((tmp_path/'parts.json').read_text())['chapter']==12
    assert parts.part_dir_name('book',12,4)=='book_12-4'


def test_parts_model_sees_causal_tail_after_eighty_characters():
    tail='他终于回答了之前的问题，两人决定一起离开。'
    prompt=_prompt(['他说：“'+'话'*100+'。”'+tail,'第二天到了实验室。'],2,[(1,1)],[])
    assert tail in prompt


def test_episode_duration_check_rejects_balanced_but_overlong_parts():
    rows=['他说：“'+'话'*300+'。”']*6
    cut=parts.mechanical_split(rows,3)
    assert parts.validate(cut,rows)==[]
    assert any('超过 120 秒' in issue for issue in parts.validate(cut,rows,max_seconds=120))


def test_parts_retry_includes_previous_cut_to_preserve_valid_boundaries():
    previous=[{'first':1,'last':2,'reason':'保留的完整交锋'}]
    prompt=_prompt(['开场。','结束。'],1,[],['第1集超时'],previous)
    assert '保留的完整交锋' in prompt and '上次实际方案' in prompt


def test_one_long_paragraph_does_not_publish_a_false_short_episode(tmp_path):
    import pytest
    from novel_manga.application.planning.parts import decide_parts
    with pytest.raises(ValueError,match='单段原文仍超时'):
        decide_parts(tmp_path,12,'他说：“'+'话'*1000+'。”')
    assert not list(tmp_path.rglob('parts.json'))
